"""Export adapter: convert section markdown to DOCX.

Uses python-docx to render markdown content into a Word document.
"""

import io
import logging
import re

logger = logging.getLogger(__name__)


def _parse_markdown_to_blocks(markdown: str) -> list[dict]:
    """Parse markdown into a list of typed blocks for rendering."""
    blocks: list[dict] = []
    for line in markdown.split("\n"):
        stripped = line.strip()
        if not stripped:
            continue

        # Headings
        if stripped.startswith("### "):
            blocks.append({"type": "h3", "text": stripped[4:]})
        elif stripped.startswith("## "):
            blocks.append({"type": "h2", "text": stripped[3:]})
        elif stripped.startswith("# "):
            blocks.append({"type": "h1", "text": stripped[2:]})
        # Bullet lists
        elif stripped.startswith("- ") or stripped.startswith("* "):
            blocks.append({"type": "bullet", "text": stripped[2:]})
        elif re.match(r"^\d+\.\s", stripped):
            blocks.append({"type": "numbered", "text": re.sub(r"^\d+\.\s", "", stripped)})
        # Bold/italic text → paragraph
        else:
            blocks.append({"type": "paragraph", "text": stripped})

    return blocks


def render_markdown_to_docx(sections: list[dict]) -> bytes:
    """Render a list of sections to a DOCX file.

    Each section dict should have:
      - title: str
      - content_markdown: str
    Returns the DOCX file as bytes.
    """
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()

    # Set default style
    style = doc.styles["Normal"]
    style.font.size = Pt(11)

    for section in sections:
        title = section.get("title", "Untitled")
        markdown_content = section.get("content_markdown", "")

        # Add section heading
        doc.add_heading(title, level=2)

        # Parse and render markdown blocks
        blocks = _parse_markdown_to_blocks(markdown_content)
        for block in blocks:
            block_type = block["type"]
            text = block["text"]

            # Strip markdown formatting markers for DOCX
            clean_text = re.sub(r"\*{1,2}(.*?)\*{1,2}", r"\1", text)
            clean_text = re.sub(r"_{1,2}(.*?)_{1,2}", r"\1", clean_text)

            if block_type == "h1":
                doc.add_heading(clean_text, level=1)
            elif block_type == "h2":
                doc.add_heading(clean_text, level=2)
            elif block_type == "h3":
                doc.add_heading(clean_text, level=3)
            elif block_type == "bullet":
                doc.add_paragraph(clean_text, style="List Bullet")
            elif block_type == "numbered":
                doc.add_paragraph(clean_text, style="List Number")
            else:
                doc.add_paragraph(clean_text)

        # Add spacing between sections
        doc.add_paragraph()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
