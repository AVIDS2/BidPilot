"""Tests for worker adapters: parser, embedding, LLM, requirements, export."""

import os
import tempfile
import io

import pytest

from app.adapters.parser import (
    _extract_csv_text,
    _extract_docx_text,
    _extract_text,
    _extract_xlsx_text,
    _split_into_chunks,
)
from app.adapters import parser as parser_adapter
from app.adapters import embedding as embedding_adapter
from app.adapters import llm as llm_adapter
from app.adapters import anthropic_llm as anthropic_llm_adapter
from app.adapters import requirements as requirements_adapter
from app.adapters.embedding import generate_embedding, generate_embeddings_batch
from contracts import EmbeddingOutcomeStatus
from app.adapters.llm import draft_section
from app.adapters.provider_errors import ProviderInvocationError, provider_error_for_status
from app.adapters.requirements import extract_requirements
from app.adapters.export import render_markdown_to_docx, _parse_markdown_to_blocks


_EMBEDDING_ENV_NAMES = (
    "DOCPILOT_EMBEDDING_API_KEY",
    "DOCPILOT_EMBEDDING_BASE_URL",
    "DOCPILOT_EMBEDDING_MODEL",
    "DOCPILOT_EMBEDDING_DIMENSIONS",
    "EMBEDDING_API_KEY",
    "EMBEDDING_API_URL",
    "EMBEDDING_MODEL",
    "EMBEDDING_DIMENSIONS",
    "OPENROUTER_API_KEY",
    "OPENROUTER_BASE_URL",
    "OPENROUTER_EMBEDDING_MODEL",
    "OPENROUTER_EMBEDDING_DIMENSIONS",
    "OPENAI_API_KEY",
    "DOCPILOT_PROVIDER_OPENAI_API_KEY",
    "DOCPILOT_PROVIDER_DOMESTIC_API_KEY",
    "DOCPILOT_PROVIDER_DOMESTIC_BASE_URL",
    "DOCPILOT_EMBEDDING_MODEL_TEXT",
    "ALIYUN_API_KEY",
    "DASHSCOPE_API_KEY",
)


def _clear_embedding_env(monkeypatch) -> None:
    for name in _EMBEDDING_ENV_NAMES:
        monkeypatch.delenv(name, raising=False)


_CHAT_ENV_NAMES = (
    "LLM_API_KEY",
    "LLM_API_URL",
    "LLM_MODEL",
    "OPENAI_API_KEY",
    "DOCPILOT_PROVIDER_OPENAI_API_KEY",
    "DOCPILOT_PROVIDER_OPENAI_BASE_URL",
    "DOCPILOT_PROVIDER_DOMESTIC_API_KEY",
    "DOCPILOT_PROVIDER_DOMESTIC_BASE_URL",
    "DOCPILOT_PROVIDER_DOMESTIC_MODEL",
    "DOCPILOT_LLM_MODEL_PRIMARY",
    "DEEPSEEK_API_KEY",
    "DEEPSEEK_BASE_URL",
    "DEEPSEEK_MODEL",
    "OPENCODE_API_KEY",
    "OPENCODE_BASE_URL",
    "OPENCODE_MODEL",
    "DOCPILOT_ASSISTANT_API_KEY",
    "DOCPILOT_ASSISTANT_PROVIDER_ID",
    "DOCPILOT_ASSISTANT_BASE_URL",
    "DOCPILOT_ASSISTANT_MODEL",
    "MIMO_API_KEY",
    "MIMO_BASE_URL",
    "MIMO_MODEL",
    "XIAOMI_API_KEY",
    "ALIYUN_API_KEY",
    "DASHSCOPE_API_KEY",
)


def _clear_chat_env(monkeypatch) -> None:
    for name in _CHAT_ENV_NAMES:
        monkeypatch.delenv(name, raising=False)


class TestParserChunking:
    def test_split_empty_text(self) -> None:
        assert _split_into_chunks("") == []
        assert _split_into_chunks("   ") == []

    def test_split_short_text_returns_single_chunk(self) -> None:
        text = "Short paragraph."
        chunks = _split_into_chunks(text)
        assert len(chunks) == 1
        content, meta = chunks[0]
        assert content == "Short paragraph."
        assert meta["chunk_type"] == "paragraphs"
        assert isinstance(meta["heading_path"], list)

    def test_split_heading_boundaries(self) -> None:
        text = "# Section 1\n\nPara one.\n\n## Sub 1.1\n\nPara two."
        chunks = _split_into_chunks(text)
        assert len(chunks) >= 2
        # Each chunk should have heading_path reflecting its section
        paths = [meta["heading_path"] for _, meta in chunks]
        assert any("Section 1" in p for p in paths)
        assert any("Sub 1.1" in p for p in paths)

    def test_table_atomic(self) -> None:
        text = "# Data\n\n| Name | Value |\n|------|-------|\n| A    | 1     |\n| B    | 2     |"
        chunks = _split_into_chunks(text)
        table_chunks = [(c, m) for c, m in chunks if m["chunk_type"] == "table"]
        assert len(table_chunks) >= 1
        _, meta = table_chunks[0]
        assert meta["table_headers"] == ["Name", "Value"]
        assert meta["table_row_count"] == 2

    def test_heading_path_tracking(self) -> None:
        text = "# Level 1\n\n## Level 2\n\n### Level 3\n\nContent here."
        chunks = _split_into_chunks(text)
        assert len(chunks) >= 1
        _, meta = chunks[-1]
        assert "Level 3" in meta["heading_path"]
        assert "Level 2" in meta["heading_path"]
        assert "Level 1" in meta["heading_path"]

    def test_long_text_multiple_chunks(self) -> None:
        paragraphs = [f"Paragraph {i} with some content here for testing purposes." for i in range(50)]
        text = "\n\n".join(paragraphs)
        chunks = _split_into_chunks(text, size=200)
        assert len(chunks) > 1

    def test_extract_text_plain_file(self) -> None:
        with tempfile.NamedTemporaryFile(mode="w", suffix=".txt", delete=False, encoding="utf-8") as f:
            f.write("Hello world from plain text file.")
            path = f.name
        try:
            result = _extract_text(path, "text/plain")
            assert "Hello world" in result
        finally:
            os.unlink(path)

    def test_docx_tables_are_preserved_as_retrievable_markdown(self) -> None:
        from docx import Document

        document = Document()
        document.add_heading("能力矩阵", level=1)
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "能力"
        table.cell(0, 1).text = "证据"
        table.cell(1, 0).text = "交付"
        table.cell(1, 1).text = "案例 A"
        output = io.BytesIO()
        document.save(output)

        text = _extract_docx_text(output.getvalue())
        assert "## Table 1" in text
        assert "| 能力 | 证据 |" in text
        assert "| 交付 | 案例 A |" in text

    def test_xlsx_and_csv_are_normalized_to_sheet_tables(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "资质"
        sheet.append(["名称", "状态"])
        sheet.append(["ISO", "有效"])
        output = io.BytesIO()
        workbook.save(output)

        xlsx_text = _extract_xlsx_text(output.getvalue())
        csv_text = _extract_csv_text("名称,状态\nISO,有效\n".encode("utf-8"))
        assert "## Sheet: 资质" in xlsx_text
        assert "| 名称 | 状态 |" in xlsx_text
        assert "## Sheet: CSV" in csv_text
        assert "| ISO | 有效 |" in csv_text

    def test_page_and_sheet_headings_become_unified_locators(self) -> None:
        chunks = _split_into_chunks("## Page 3\n\n扫描页内容")
        assert chunks[0][1]["source_locator"] == {"page": 3}

    def test_scanned_pdf_without_ocr_is_a_non_retryable_parse_failure(self, monkeypatch) -> None:
        monkeypatch.setattr(parser_adapter, "_download_from_minio", lambda _key: b"%PDF-scan")
        monkeypatch.setattr(
            parser_adapter,
            "_extract_pdf_outcome",
            lambda _data: parser_adapter._ExtractionOutcome(
                text="",
                error_code="pdf_ocr_unavailable",
                retryable=False,
            ),
        )

        result = _extract_text("private/scanned.pdf", "application/pdf")

        assert isinstance(result, str)
        assert result == ""
        assert result.error_code == "pdf_ocr_unavailable"
        assert result.retryable is False


class TestEmbeddingAdapter:
    def test_embedding_without_api_key_is_explicitly_not_configured(self, monkeypatch) -> None:
        _clear_embedding_env(monkeypatch)
        result = generate_embedding("test text")
        assert result.model == "not_configured"
        assert result.status is EmbeddingOutcomeStatus.NOT_CONFIGURED
        assert result.embedding is None
        assert result.error_code == "embedding_not_configured"

    def test_batch_without_api_key_is_not_configured_without_zero_vectors(self, monkeypatch) -> None:
        _clear_embedding_env(monkeypatch)
        results = generate_embeddings_batch(["text1", "text2"])
        assert len(results) == 2
        assert all(result.status is EmbeddingOutcomeStatus.NOT_CONFIGURED for result in results)
        assert all(result.embedding is None for result in results)

    def test_large_batch_is_split_into_bounded_provider_requests(self, monkeypatch) -> None:
        calls: list[list[str]] = []

        def _fake_request(input_value: str | list[str], *, expected_count: int):
            assert isinstance(input_value, list)
            calls.append(input_value)
            return [
                embedding_adapter.EmbeddingResult(
                    status=EmbeddingOutcomeStatus.SUCCESS,
                    model=value,
                    profile_id="test:embedding:1536:bidpilot-lexical-v1",
                    vector=[0.1] * 1536,
                )
                for value in input_value
            ]

        monkeypatch.setattr(embedding_adapter, "_request_embeddings", _fake_request)

        results = generate_embeddings_batch([f"chunk-{index}" for index in range(9)])

        assert calls == [
            ["chunk-0", "chunk-1", "chunk-2", "chunk-3"],
            ["chunk-4", "chunk-5", "chunk-6", "chunk-7"],
            ["chunk-8"],
        ]
        assert [result.model for result in results] == [f"chunk-{index}" for index in range(9)]

    def test_domestic_embedding_env_uses_dashscope_defaults(self, monkeypatch) -> None:
        monkeypatch.delenv("EMBEDDING_API_KEY", raising=False)
        monkeypatch.delenv("EMBEDDING_API_URL", raising=False)
        monkeypatch.delenv("EMBEDDING_MODEL", raising=False)
        monkeypatch.delenv("EMBEDDING_DIMENSIONS", raising=False)
        monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)
        monkeypatch.delenv("OPENROUTER_BASE_URL", raising=False)
        monkeypatch.delenv("OPENROUTER_EMBEDDING_MODEL", raising=False)
        monkeypatch.delenv("OPENROUTER_EMBEDDING_DIMENSIONS", raising=False)
        monkeypatch.delenv("OPENAI_API_KEY", raising=False)
        monkeypatch.delenv("DOCPILOT_PROVIDER_OPENAI_API_KEY", raising=False)
        monkeypatch.delenv("DOCPILOT_PROVIDER_DOMESTIC_BASE_URL", raising=False)
        monkeypatch.delenv("DOCPILOT_PROVIDER_DOMESTIC_API_KEY", raising=False)
        monkeypatch.delenv("DOCPILOT_EMBEDDING_MODEL_TEXT", raising=False)
        monkeypatch.delenv("DASHSCOPE_API_KEY", raising=False)
        monkeypatch.setenv("ALIYUN_API_KEY", "test-aliyun-key")

        assert embedding_adapter._api_key() == "test-aliyun-key"
        assert embedding_adapter._api_url() == "https://dashscope.aliyuncs.com/compatible-mode/v1/embeddings"
        assert embedding_adapter._api_model() == "text-embedding-v4"

    def test_openrouter_embedding_env_uses_qwen_defaults(self, monkeypatch) -> None:
        monkeypatch.delenv("EMBEDDING_API_KEY", raising=False)
        monkeypatch.delenv("EMBEDDING_API_URL", raising=False)
        monkeypatch.delenv("EMBEDDING_MODEL", raising=False)
        monkeypatch.delenv("EMBEDDING_DIMENSIONS", raising=False)
        monkeypatch.delenv("OPENROUTER_BASE_URL", raising=False)
        monkeypatch.delenv("OPENROUTER_EMBEDDING_MODEL", raising=False)
        monkeypatch.delenv("OPENROUTER_EMBEDDING_DIMENSIONS", raising=False)
        monkeypatch.delenv("OPENAI_API_KEY", raising=False)
        monkeypatch.delenv("DOCPILOT_PROVIDER_OPENAI_API_KEY", raising=False)
        monkeypatch.delenv("DOCPILOT_PROVIDER_DOMESTIC_BASE_URL", raising=False)
        monkeypatch.delenv("DOCPILOT_PROVIDER_DOMESTIC_API_KEY", raising=False)
        monkeypatch.delenv("DOCPILOT_EMBEDDING_MODEL_TEXT", raising=False)
        monkeypatch.delenv("ALIYUN_API_KEY", raising=False)
        monkeypatch.delenv("DASHSCOPE_API_KEY", raising=False)
        monkeypatch.setenv("OPENROUTER_API_KEY", "test-openrouter-key")

        assert embedding_adapter._api_key() == "test-openrouter-key"
        assert embedding_adapter._api_url() == "https://openrouter.ai/api/v1/embeddings"
        assert embedding_adapter._api_model() == "qwen/qwen3-embedding-8b"
        assert embedding_adapter._api_dimensions() == 1536

    def test_explicit_platform_embedding_env_has_priority(self, monkeypatch) -> None:
        _clear_embedding_env(monkeypatch)
        monkeypatch.setenv("DOCPILOT_EMBEDDING_API_KEY", "test-platform-key")
        monkeypatch.setenv("DOCPILOT_EMBEDDING_BASE_URL", "https://embeddings.example.test/v1")
        monkeypatch.setenv("DOCPILOT_EMBEDDING_MODEL", "qwen/qwen3-embedding-8b")
        monkeypatch.setenv("DOCPILOT_EMBEDDING_DIMENSIONS", "1536")

        assert embedding_adapter._api_key() == "test-platform-key"
        assert embedding_adapter._api_url() == "https://embeddings.example.test/v1/embeddings"
        assert embedding_adapter._api_model() == "qwen/qwen3-embedding-8b"
        assert embedding_adapter._api_dimensions() == 1536

    def test_openrouter_embedding_request_includes_dimensions(self, monkeypatch) -> None:
        captured: dict[str, object] = {}

        class Response:
            def raise_for_status(self) -> None:
                return None

            def json(self) -> dict[str, object]:
                return {
                    "data": [{"embedding": [0.1] * 1536}],
                    "usage": {"total_tokens": 3},
                }

        def fake_post(url: str, *, headers: dict[str, str], json: dict[str, object], timeout: float) -> Response:
            captured["url"] = url
            captured["headers"] = headers
            captured["json"] = json
            captured["timeout"] = timeout
            return Response()

        monkeypatch.delenv("EMBEDDING_API_KEY", raising=False)
        monkeypatch.delenv("EMBEDDING_API_URL", raising=False)
        monkeypatch.delenv("EMBEDDING_MODEL", raising=False)
        monkeypatch.delenv("EMBEDDING_DIMENSIONS", raising=False)
        monkeypatch.delenv("OPENROUTER_BASE_URL", raising=False)
        monkeypatch.delenv("OPENROUTER_EMBEDDING_MODEL", raising=False)
        monkeypatch.delenv("OPENROUTER_EMBEDDING_DIMENSIONS", raising=False)
        monkeypatch.delenv("OPENAI_API_KEY", raising=False)
        monkeypatch.delenv("DOCPILOT_PROVIDER_OPENAI_API_KEY", raising=False)
        monkeypatch.delenv("DOCPILOT_PROVIDER_DOMESTIC_API_KEY", raising=False)
        monkeypatch.delenv("ALIYUN_API_KEY", raising=False)
        monkeypatch.delenv("DASHSCOPE_API_KEY", raising=False)
        monkeypatch.setenv("OPENROUTER_API_KEY", "test-openrouter-key")
        monkeypatch.setattr(embedding_adapter.httpx, "post", fake_post)

        result = generate_embedding("hello")

        assert result.model == "qwen/qwen3-embedding-8b"
        assert result.status is EmbeddingOutcomeStatus.SUCCESS
        assert len(result.embedding) == 1536
        assert result.profile_id == "openrouter:qwen/qwen3-embedding-8b:1536:bidpilot-lexical-v1"
        assert result.token_count == 3
        assert result.usage_reported is True
        assert captured["url"] == "https://openrouter.ai/api/v1/embeddings"
        assert captured["json"] == {
            "input": "hello",
            "model": "qwen/qwen3-embedding-8b",
            "dimensions": 1536,
        }

    def test_timeout_returns_retryable_outcome_without_vector(self, monkeypatch) -> None:
        _clear_embedding_env(monkeypatch)
        monkeypatch.setenv("OPENROUTER_API_KEY", "test-openrouter-key")

        def raise_timeout(*_args, **_kwargs):
            raise embedding_adapter.httpx.TimeoutException("simulated timeout")

        monkeypatch.setattr(embedding_adapter.httpx, "post", raise_timeout)

        result = generate_embedding("hello")

        assert result.status is EmbeddingOutcomeStatus.TRANSIENT_FAILURE
        assert result.embedding is None
        assert result.error_code == "provider_timeout"

    def test_dimension_mismatch_returns_explicit_outcome_without_vector(self, monkeypatch) -> None:
        _clear_embedding_env(monkeypatch)
        monkeypatch.setenv("OPENROUTER_API_KEY", "test-openrouter-key")

        class Response:
            def raise_for_status(self) -> None:
                return None

            def json(self) -> dict[str, object]:
                return {"data": [{"embedding": [0.1] * 1024}]}

        monkeypatch.setattr(embedding_adapter.httpx, "post", lambda *_args, **_kwargs: Response())

        result = generate_embedding("hello")

        assert result.status is EmbeddingOutcomeStatus.DIMENSION_MISMATCH
        assert result.embedding is None
        assert result.error_code == "embedding_dimension_mismatch"


class TestLLMAdapter:
    def test_stub_draft_without_api_key(self) -> None:
        os.environ.pop("LLM_API_KEY", None)
        os.environ.pop("OPENAI_API_KEY", None)
        os.environ.pop("DOCPILOT_PROVIDER_OPENAI_API_KEY", None)
        os.environ.pop("DOCPILOT_PROVIDER_DOMESTIC_API_KEY", None)
        os.environ.pop("ALIYUN_API_KEY", None)
        os.environ.pop("DASHSCOPE_API_KEY", None)
        result = draft_section("technical-approach", ["evidence text"], "proj1")
        assert result.model_used == "stub"
        assert "technical-approach" in result.content_markdown.lower() or "Technical Approach" in result.content_markdown
        assert "evidence text" in result.content_markdown

    def test_domestic_llm_env_uses_dashscope_defaults(self, monkeypatch) -> None:
        monkeypatch.delenv("LLM_API_KEY", raising=False)
        monkeypatch.delenv("LLM_API_URL", raising=False)
        monkeypatch.delenv("LLM_MODEL", raising=False)
        monkeypatch.delenv("OPENAI_API_KEY", raising=False)
        monkeypatch.delenv("DOCPILOT_PROVIDER_OPENAI_API_KEY", raising=False)
        monkeypatch.delenv("DOCPILOT_PROVIDER_DOMESTIC_BASE_URL", raising=False)
        monkeypatch.delenv("DOCPILOT_PROVIDER_DOMESTIC_API_KEY", raising=False)
        monkeypatch.delenv("DOCPILOT_LLM_MODEL_PRIMARY", raising=False)
        monkeypatch.delenv("ALIYUN_API_KEY", raising=False)
        monkeypatch.setenv("DASHSCOPE_API_KEY", "test-dashscope-key")

        assert llm_adapter._api_key() == "test-dashscope-key"
        assert llm_adapter._api_url() == "https://dashscope.aliyuncs.com/compatible-mode/v1/chat/completions"
        assert llm_adapter._api_model() == "qwen3.5-flash"

    def test_deepseek_llm_env_uses_chat_completions_and_supported_default(self, monkeypatch) -> None:
        _clear_chat_env(monkeypatch)
        monkeypatch.setenv("DEEPSEEK_API_KEY", "test-deepseek-key")

        assert llm_adapter._api_key() == "test-deepseek-key"
        assert llm_adapter._api_url() == "https://api.deepseek.com/v1/chat/completions"
        assert llm_adapter._api_model() == "deepseek-v4-flash"

    def test_opencode_go_llm_env_is_preferred_over_legacy_deepseek(self, monkeypatch) -> None:
        _clear_chat_env(monkeypatch)
        monkeypatch.setenv("OPENCODE_API_KEY", "test-opencode-key")
        monkeypatch.setenv("DEEPSEEK_API_KEY", "test-deepseek-key")

        assert llm_adapter._api_key() == "test-opencode-key"
        assert llm_adapter._api_url() == "https://opencode.ai/zen/go/v1/chat/completions"
        assert llm_adapter._api_model() == "deepseek-v4-flash"

    def test_opencode_go_is_not_redirected_by_legacy_endpoint_variables(self, monkeypatch) -> None:
        _clear_chat_env(monkeypatch)
        monkeypatch.setenv("OPENCODE_API_KEY", "test-opencode-key")
        monkeypatch.setenv("LLM_API_KEY", "legacy-key")
        monkeypatch.setenv("LLM_API_URL", "https://legacy.example.test/v1/chat/completions")
        monkeypatch.setenv("LLM_MODEL", "legacy-model")

        assert llm_adapter._api_key() == "test-opencode-key"
        assert llm_adapter._api_url() == "https://opencode.ai/zen/go/v1/chat/completions"
        assert llm_adapter._api_model() == "deepseek-v4-flash"

    def test_opencode_go_uses_defaults_when_compose_injects_blank_overrides(self, monkeypatch) -> None:
        _clear_chat_env(monkeypatch)
        monkeypatch.setenv("OPENCODE_API_KEY", "test-opencode-key")
        monkeypatch.setenv("OPENCODE_BASE_URL", "")
        monkeypatch.setenv("OPENCODE_MODEL", "")

        assert llm_adapter._api_url() == "https://opencode.ai/zen/go/v1/chat/completions"
        assert llm_adapter._api_model() == "deepseek-v4-flash"

    def test_mimo_direct_balance_env_uses_official_endpoint_and_model(self, monkeypatch) -> None:
        _clear_chat_env(monkeypatch)
        monkeypatch.setenv("DOCPILOT_ASSISTANT_API_KEY", "test-mimo-key")
        monkeypatch.setenv("DOCPILOT_ASSISTANT_PROVIDER_ID", "mimo")

        assert llm_adapter._api_key() == "test-mimo-key"
        assert llm_adapter._api_url() == "https://api.xiaomimimo.com/v1/chat/completions"
        assert llm_adapter._api_model() == "mimo-v2.5-pro"

    def test_timeout_raises_a_retryable_provider_error(self, monkeypatch) -> None:
        _clear_chat_env(monkeypatch)
        monkeypatch.setenv("LLM_API_KEY", "test-provider-key")

        def raise_timeout(*_args, **_kwargs):
            raise llm_adapter.httpx.TimeoutException("simulated timeout")

        monkeypatch.setattr(llm_adapter.httpx, "post", raise_timeout)

        with pytest.raises(ProviderInvocationError) as error:
            draft_section("technical-approach", ["evidence"], "project-1")

        assert error.value.error_code == "provider_timeout"
        assert error.value.retryable is True

    def test_truncated_reasoning_response_is_not_blindly_retried(self, monkeypatch) -> None:
        _clear_chat_env(monkeypatch)
        monkeypatch.setenv("LLM_API_KEY", "test-provider-key")
        captured: dict[str, object] = {}

        class Response:
            status_code = 200

            def json(self) -> dict[str, object]:
                return {
                    "choices": [
                        {
                            "finish_reason": "length",
                            "message": {
                                "content": "",
                                "reasoning_content": "internal reasoning only",
                            },
                        }
                    ]
                }

        def post(*_args, **kwargs):
            captured.update(kwargs["json"])
            return Response()

        monkeypatch.setattr(llm_adapter.httpx, "post", post)

        with pytest.raises(ProviderInvocationError) as error:
            draft_section("technical-approach", ["evidence"], "project-1")

        assert error.value.error_code == "provider_response_truncated"
        assert error.value.retryable is False
        assert captured["max_tokens"] == llm_adapter._MAX_DRAFT_OUTPUT_TOKENS

    def test_draft_budget_supports_reasoning_and_full_section_output(self) -> None:
        assert llm_adapter._MAX_DRAFT_OUTPUT_TOKENS == 16_000

    def test_production_missing_provider_never_returns_stub(self, monkeypatch) -> None:
        _clear_chat_env(monkeypatch)
        monkeypatch.setenv("DOCPILOT_ENV", "production")
        monkeypatch.delenv("DOCPILOT_ALLOW_STUB_LLM", raising=False)

        with pytest.raises(ProviderInvocationError) as error:
            draft_section("technical-approach", ["evidence"], "project-1")

        assert error.value.error_code == "provider_not_configured"
        assert error.value.retryable is False

    def test_anthropic_auth_failure_is_not_retryable(self, monkeypatch) -> None:
        monkeypatch.setenv("ANTHROPIC_API_KEY", "test-provider-key")

        class Response:
            status_code = 401

        monkeypatch.setattr(anthropic_llm_adapter.httpx, "post", lambda *_args, **_kwargs: Response())

        with pytest.raises(ProviderInvocationError) as error:
            anthropic_llm_adapter.draft_section("technical-approach", ["evidence"], "project-1")

        assert error.value.error_code == "provider_auth_failed"
        assert error.value.retryable is False


def test_provider_status_classification_is_safe_and_deterministic() -> None:
    rate_limited = provider_error_for_status(429)
    invalid = provider_error_for_status(422)

    assert rate_limited.error_code == "provider_rate_limited"
    assert rate_limited.retryable is True
    assert invalid.error_code == "provider_request_invalid"
    assert invalid.retryable is False


class TestProviderEnv:
    def test_requirements_adapter_uses_domestic_chat_config(self, monkeypatch) -> None:
        monkeypatch.delenv("LLM_API_KEY", raising=False)
        monkeypatch.delenv("LLM_API_URL", raising=False)
        monkeypatch.delenv("LLM_MODEL", raising=False)
        monkeypatch.delenv("OPENAI_API_KEY", raising=False)
        monkeypatch.delenv("DOCPILOT_PROVIDER_OPENAI_API_KEY", raising=False)
        monkeypatch.delenv("ALIYUN_API_KEY", raising=False)
        monkeypatch.delenv("DASHSCOPE_API_KEY", raising=False)
        monkeypatch.setenv("DOCPILOT_PROVIDER_DOMESTIC_API_KEY", "test-domestic-key")
        monkeypatch.setenv("DOCPILOT_PROVIDER_DOMESTIC_BASE_URL", "https://example.test/v1/")
        monkeypatch.setenv("DOCPILOT_LLM_MODEL_PRIMARY", "qwen-test")

        assert requirements_adapter._api_key() == "test-domestic-key"
        assert requirements_adapter._api_url() == "https://example.test/v1/chat/completions"
        assert requirements_adapter._api_model() == "qwen-test"


class TestRequirementsAdapter:
    def test_stub_extracts_shall_must(self) -> None:
        os.environ.pop("LLM_API_KEY", None)
        os.environ.pop("OPENAI_API_KEY", None)
        chunks = [
            "The system shall provide user authentication.\nUsers must change passwords every 90 days.",
            "The interface should be responsive.",
            "This is just a description with no requirements.",
        ]
        results = extract_requirements(chunks, "proj1")
        assert len(results) >= 2
        assert any("shall" in r.requirement_text.lower() or "must" in r.requirement_text.lower() for r in results)
        # High priority for shall/must
        high = [r for r in results if r.priority == "high"]
        assert len(high) >= 1


class TestExportAdapter:
    def test_parse_markdown_blocks(self) -> None:
        md = "# Title\n\n## Subtitle\n\nParagraph text\n\n- Bullet 1\n- Bullet 2\n\n1. Numbered item"
        blocks = _parse_markdown_to_blocks(md)
        types = [b["type"] for b in blocks]
        assert "h1" in types
        assert "h2" in types
        assert "paragraph" in types
        assert "bullet" in types
        assert "numbered" in types

    def test_render_docx_returns_bytes(self) -> None:
        sections = [
            {"title": "Section 1", "content_markdown": "## Heading\n\nSome text here.\n\n- Item A\n- Item B"},
            {"title": "Section 2", "content_markdown": "Another section with **bold** text."},
        ]
        result = render_markdown_to_docx(sections)
        assert isinstance(result, bytes)
        assert len(result) > 100  # DOCX should be non-trivial
        # Verify it's a valid ZIP (DOCX is a ZIP archive)
        import zipfile
        buf = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
        buf.write(result)
        buf.close()
        try:
            assert zipfile.is_zipfile(buf.name)
        finally:
            os.unlink(buf.name)
