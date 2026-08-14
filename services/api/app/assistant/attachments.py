"""Assistant attachment extraction helpers."""

from __future__ import annotations

from datetime import UTC, datetime, timedelta
from hashlib import sha256
from io import BytesIO
import logging
import os
from pathlib import Path
import shutil
import subprocess
from time import time
from uuid import uuid4

from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.access.service import require_project_capability
from app.audit.service import record_audit_event
from app.auth.schemas import CurrentUser
from app.celery_client import celery
from app.models import AssistantAttachment, Bundle, SourceDocument
from app.usage.schemas import ProviderSource
from app.usage.service import EMBEDDING_INDEX_STARTED, check_indexing_quota, record_usage_event

from .schemas import (
    AssistantAttachmentKind,
    AssistantAttachmentPayload,
    AssistantAttachmentUploadResponse,
)

MAX_ATTACHMENT_BYTES = 10 * 1024 * 1024
MAX_ATTACHMENT_TEXT_CHARS = 16_000
MAX_TOTAL_ATTACHMENT_TEXT_CHARS = 24_000
ATTACHMENT_CACHE_TTL_SECONDS = 60 * 60
ASSISTANT_ATTACHMENT_RETENTION = timedelta(hours=24)
MAX_STAGED_ATTACHMENTS_PER_ACTION = 6

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".json",
    ".jsonl",
    ".yaml",
    ".yml",
    ".xml",
    ".html",
    ".htm",
    ".log",
}

_ATTACHMENT_TEXT_CACHE: dict[str, tuple[float, str]] = {}
logger = logging.getLogger(__name__)


def stage_assistant_attachment(
    db: Session,
    *,
    current_user: CurrentUser,
    filename: str,
    content_type: str,
    data: bytes,
    kind: AssistantAttachmentKind,
    extraction: AssistantAttachmentUploadResponse,
) -> AssistantAttachmentUploadResponse:
    """Persist an uploaded file before it becomes project evidence.

    The raw object is private to the uploading user and organization.  It is
    only copied into a project after the governed ingestion capability runs.
    """
    now = _utcnow()
    attachment = AssistantAttachment(
        user_id=current_user.id,
        org_id=current_user.org_id,
        storage_key="pending",
        original_filename=filename or "untitled",
        mime_type=content_type or "application/octet-stream",
        kind=kind,
        size=len(data),
        checksum=sha256(data).hexdigest(),
        extraction_status=extraction.extraction_status,
        extracted_text=extraction.extracted_text,
        extraction_error=extraction.error,
        status="staging",
        expires_at=now + ASSISTANT_ATTACHMENT_RETENTION,
    )
    db.add(attachment)
    db.flush()

    try:
        from app.adapters.storage import upload_assistant_staging_bytes

        attachment.storage_key = upload_assistant_staging_bytes(
            org_id=current_user.org_id,
            user_id=current_user.id,
            attachment_id=attachment.id,
            data=data,
            content_type=attachment.mime_type,
        )
    except Exception as exc:
        db.rollback()
        logger.warning("Assistant attachment staging failed: %s", type(exc).__name__)
        raise HTTPException(status_code=503, detail="Unable to stage attachment") from exc

    attachment.status = "staged"
    try:
        db.commit()
    except Exception as exc:
        db.rollback()
        _delete_storage_keys_quietly([attachment.storage_key])
        logger.warning("Assistant attachment staging persistence failed: %s", type(exc).__name__)
        raise HTTPException(status_code=503, detail="Unable to stage attachment") from exc
    db.refresh(attachment)
    return _upload_response_from_record(attachment)


def hydrate_assistant_attachments(
    db: Session,
    *,
    current_user: CurrentUser,
    attachments: list[AssistantAttachmentPayload],
) -> list[AssistantAttachmentPayload]:
    """Load attachment text from the server-side record, never from the browser."""
    if not attachments:
        return []
    if len(attachments) > MAX_STAGED_ATTACHMENTS_PER_ACTION:
        raise HTTPException(status_code=400, detail="Too many attachments in one assistant turn")

    attachment_ids = [attachment.id for attachment in attachments]
    if any(not attachment_id for attachment_id in attachment_ids):
        raise HTTPException(status_code=400, detail="Assistant attachments must be uploaded before use")
    if len(set(attachment_ids)) != len(attachment_ids):
        raise HTTPException(status_code=400, detail="Duplicate assistant attachment ids are not allowed")

    records = {
        record.id: record
        for record in db.scalars(
            select(AssistantAttachment).where(
                AssistantAttachment.id.in_(attachment_ids),
                AssistantAttachment.user_id == current_user.id,
                AssistantAttachment.org_id == current_user.org_id,
            )
        )
    }
    if len(records) != len(attachment_ids):
        raise HTTPException(status_code=404, detail="Assistant attachment not found")

    hydrated: list[AssistantAttachmentPayload] = []
    for attachment_id in attachment_ids:
        record = records[attachment_id or ""]
        if record.status != "attached":
            _require_staged_attachment(record)
        hydrated.append(_payload_from_record(record))
    return hydrated


def attachment_planner_context(attachments: list[AssistantAttachmentPayload]) -> list[dict[str, str | int]]:
    """Return non-sensitive attachment metadata for the bounded operator planner."""
    return [
        {
            "id": attachment.id or "",
            "name": attachment.name[:255],
            "kind": attachment.kind,
            "mime_type": attachment.mime_type or "application/octet-stream",
            "size": attachment.size or 0,
        }
        for attachment in attachments
        if attachment.id and not attachment.document_id
    ]


def attach_staged_attachments_to_project(
    db: Session,
    *,
    current_user: CurrentUser,
    project_id: str,
    attachment_ids: list[str],
    bundle_label: str | None = None,
) -> dict[str, object]:
    """Copy approved staged attachments into a project bundle and queue ingestion."""
    if not attachment_ids:
        raise ValueError("At least one staged attachment is required")
    if len(attachment_ids) > MAX_STAGED_ATTACHMENTS_PER_ACTION:
        raise ValueError(f"At most {MAX_STAGED_ATTACHMENTS_PER_ACTION} attachments can be ingested at once")
    if len(set(attachment_ids)) != len(attachment_ids):
        raise ValueError("Duplicate staged attachment ids are not allowed")

    project = require_project_capability(
        db,
        current_user=current_user,
        project_id=project_id,
        capability="bundles.write",
    ).project
    attachments = _load_staged_attachments_for_ingestion(
        db,
        current_user=current_user,
        attachment_ids=attachment_ids,
    )
    check_indexing_quota(db, current_user.id, current_user.org_id, ProviderSource.OFFICIAL)

    label = (bundle_label or "Agent 上传资料").strip()[:255] or "Agent 上传资料"
    bundle = _find_reusable_agent_bundle(db, project.id, label)
    if bundle is None:
        bundle = Bundle(
            project_id=project.id,
            label=label,
            source_type="assistant_upload",
            ingest_status="awaiting_upload",
        )
        db.add(bundle)
        db.flush()

    copied_storage_keys: list[str] = []
    documents: list[SourceDocument] = []
    try:
        from app.adapters.storage import download_storage_key, upload_bytes

        for attachment in attachments:
            data = download_storage_key(attachment.storage_key)
            storage_key = upload_bytes(
                project.id,
                f"{bundle.id}/{uuid4().hex}-{_safe_filename(attachment.original_filename)}",
                data,
                attachment.mime_type,
            )
            copied_storage_keys.append(storage_key)
            document = SourceDocument(
                bundle_id=bundle.id,
                storage_key=storage_key,
                mime_type=attachment.mime_type,
                checksum=attachment.checksum,
                original_filename=attachment.original_filename,
                parse_status="pending",
            )
            db.add(document)
            db.flush()
            attachment.status = "attached"
            attachment.project_id = project.id
            attachment.bundle_id = bundle.id
            attachment.document_id = document.id
            attachment.attached_at = _utcnow()
            documents.append(document)

        bundle.ingest_status = "queued"
        record_usage_event(
            db,
            user_id=current_user.id,
            org_id=current_user.org_id,
            project_id=project.id,
            event_type=EMBEDDING_INDEX_STARTED,
            provider_source=ProviderSource.OFFICIAL,
            metadata_json={
                "bundle_id": bundle.id,
                "attachment_count": len(attachments),
                "action": "assistant_attachment_ingest",
            },
        )
        record_audit_event(
            db,
            project_id=project.id,
            event_type="assistant.attachments_ingested",
            actor_type="assistant",
            actor_id=current_user.id,
            payload={
                "bundle_id": bundle.id,
                "attachment_ids": [attachment.id for attachment in attachments],
                "document_ids": [document.id for document in documents],
                "attachment_count": len(documents),
            },
        )
        db.commit()
    except Exception:
        db.rollback()
        _delete_storage_keys_quietly(copied_storage_keys)
        raise

    queued = _queue_bundle_ingestion(db, bundle, current_user)
    _delete_attached_staging_objects(db, attachments)
    return {
        "project_id": project.id,
        "bundle_id": bundle.id,
        "bundle_label": bundle.label,
        "document_ids": [document.id for document in documents],
        "attachment_count": len(documents),
        "ingest_queued": queued,
    }


def link_staged_attachment_to_document(
    db: Session,
    *,
    current_user: CurrentUser,
    attachment_id: str,
    document: SourceDocument,
    project_id: str,
) -> str:
    """Mark a browser-uploaded staging record as consumed by the same document."""
    attachment = db.scalar(
        select(AssistantAttachment)
        .where(
            AssistantAttachment.id == attachment_id,
            AssistantAttachment.user_id == current_user.id,
            AssistantAttachment.org_id == current_user.org_id,
        )
        .with_for_update()
    )
    if attachment is None:
        raise HTTPException(status_code=404, detail="Assistant attachment not found")
    _require_staged_attachment(attachment)
    if attachment.checksum != document.checksum:
        raise HTTPException(status_code=409, detail="Assistant attachment does not match uploaded document")

    attachment.status = "attached"
    attachment.project_id = project_id
    attachment.bundle_id = document.bundle_id
    attachment.document_id = document.id
    attachment.attached_at = _utcnow()
    record_audit_event(
        db,
        project_id=project_id,
        event_type="assistant.attachment_linked",
        actor_type="user",
        actor_id=current_user.id,
        payload={"attachment_id": attachment.id, "document_id": document.id},
    )
    return attachment.storage_key


def delete_staged_attachment_storage(db: Session, attachment_id: str) -> bool:
    """Remove an already-consumed staging object without touching project evidence."""
    attachment = db.get(AssistantAttachment, attachment_id)
    if attachment is None or not attachment.storage_key:
        return True
    if attachment.status not in {"attached", "expired"}:
        return False
    try:
        from app.adapters.storage import delete_storage_key

        delete_storage_key(attachment.storage_key)
    except Exception as exc:
        logger.warning("Assistant attachment staged-object cleanup failed: %s", type(exc).__name__)
        return False
    attachment.storage_key = ""
    db.commit()
    return True


def extract_attachment_text(
    *,
    filename: str,
    content_type: str,
    data: bytes,
    kind: AssistantAttachmentKind | None = None,
) -> AssistantAttachmentUploadResponse:
    """Extract readable text for a chat-scoped assistant attachment."""
    safe_name = filename or "untitled"
    mime_type = content_type or "application/octet-stream"
    attachment_kind: AssistantAttachmentKind = kind or ("image" if mime_type.startswith("image/") else "file")

    if attachment_kind == "image" or mime_type.startswith("image/"):
        return _extract_image_ocr(
            name=safe_name,
            mime_type=mime_type,
            data=data,
        )

    text = ""
    error: str | None = None
    try:
        if _is_docx(safe_name, mime_type):
            text = _extract_docx(data)
        elif _is_pdf(safe_name, mime_type):
            text = _extract_pdf(data)
            if not text:
                error = "PDF 文本提取失败或文件不可选择文本。"
        elif _is_plain_text(safe_name, mime_type):
            text = _decode_text(data)
        else:
            return _response(
                name=safe_name,
                kind=attachment_kind,
                mime_type=mime_type,
                size=len(data),
                status="unsupported",
                text="",
                error="暂不支持解析该附件格式。",
            )
    except Exception as exc:  # Defensive: never let a bad attachment break chat.
        logger.warning(
            "Assistant attachment extraction failed: filename=%s error_type=%s",
            safe_name,
            type(exc).__name__,
        )
        return _response(
            name=safe_name,
            kind=attachment_kind,
            mime_type=mime_type,
            size=len(data),
            status="failed",
            text="",
            error="附件内容解析失败，请确认文件完整后重新上传。",
        )

    text = _normalize_text(text)
    if not text:
        return _response(
            name=safe_name,
            kind=attachment_kind,
            mime_type=mime_type,
            size=len(data),
            status="empty",
            text="",
            error=error,
        )

    return _response(
        name=safe_name,
        kind=attachment_kind,
        mime_type=mime_type,
        size=len(data),
        status="extracted",
        text=_truncate(text, MAX_ATTACHMENT_TEXT_CHARS),
        error=error,
    )


def build_attachment_context(message: str, attachments: list[AssistantAttachmentPayload]) -> str:
    """Append internal attachment context for the model without showing it in chat UI."""
    if not attachments:
        return message

    remaining = MAX_TOTAL_ATTACHMENT_TEXT_CHARS
    sections = [
        "【系统附件上下文】",
        "以下内容来自用户本轮上传的附件。优先结合可读附件正文回答；不要声称读取了未提供正文的图片、截图或不支持文件。",
    ]

    for index, attachment in enumerate(attachments, start=1):
        name = attachment.name
        mime_type = attachment.mime_type or "unknown"
        size = attachment.size if attachment.size is not None else 0
        header = f"附件 {index}: {name} ({attachment.kind}, {mime_type}, {size} bytes)"
        text = (attachment.extracted_text or _resolve_cached_attachment_text(attachment.id)).strip()

        if text and remaining > 0:
            clipped = _truncate(text, min(MAX_ATTACHMENT_TEXT_CHARS, remaining))
            remaining -= len(clipped)
            sections.append(f"{header}\n可读正文：\n{clipped}")
            continue

        if attachment.kind == "image":
            sections.append(
                f"{header}\n说明：当前不能读取图片像素或 OCR；图片没有可用的 OCR 正文。"
                f"{attachment.error or '请结合用户对图片的描述。'}"
            )
            continue

        reason = attachment.error or _status_reason(attachment.extraction_status)
        sections.append(f"{header}\n说明：未获得可读正文。{reason}")

    return "\n\n".join([message, "\n\n".join(sections)]).strip()


def remember_attachment_text(extraction: AssistantAttachmentUploadResponse) -> None:
    """Keep extracted text server-side for the next assistant turn."""
    text = extraction.extracted_text.strip()
    if not text:
        return
    _purge_expired_attachment_text()
    _ATTACHMENT_TEXT_CACHE[extraction.id] = (time(), text)


def _response(
    *,
    name: str,
    kind: AssistantAttachmentKind,
    mime_type: str,
    size: int,
    status: str,
    text: str,
    error: str | None = None,
) -> AssistantAttachmentUploadResponse:
    return AssistantAttachmentUploadResponse(
        id=f"att_{uuid4().hex}",
        name=name,
        kind=kind,
        mime_type=mime_type,
        size=size,
        extraction_status=status,  # type: ignore[arg-type]
        extracted_text=text,
        error=error,
    )


def _resolve_cached_attachment_text(attachment_id: str | None) -> str:
    if not attachment_id:
        return ""
    _purge_expired_attachment_text()
    cached = _ATTACHMENT_TEXT_CACHE.get(attachment_id)
    return cached[1] if cached else ""


def _purge_expired_attachment_text() -> None:
    now = time()
    expired = [
        attachment_id
        for attachment_id, (created_at, _text) in _ATTACHMENT_TEXT_CACHE.items()
        if now - created_at > ATTACHMENT_CACHE_TTL_SECONDS
    ]
    for attachment_id in expired:
        _ATTACHMENT_TEXT_CACHE.pop(attachment_id, None)


def _is_docx(filename: str, mime_type: str) -> bool:
    return Path(filename).suffix.lower() == ".docx" or mime_type == _DOCX_MIME


def _is_pdf(filename: str, mime_type: str) -> bool:
    return Path(filename).suffix.lower() == ".pdf" or mime_type == "application/pdf"


def _is_plain_text(filename: str, mime_type: str) -> bool:
    return mime_type.startswith("text/") or Path(filename).suffix.lower() in _TEXT_EXTENSIONS


def _extract_docx(data: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(data))
    parts: list[str] = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _extract_pdf(data: bytes) -> str:
    try:
        import pymupdf  # type: ignore[import-not-found]
    except Exception:
        return ""

    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        return "\n\n".join(
            doc.load_page(page_number).get_text()
            for page_number in range(doc.page_count)
        )
    finally:
        doc.close()


def _extract_image_ocr(*, name: str, mime_type: str, data: bytes) -> AssistantAttachmentUploadResponse:
    """Run bounded local OCR before the assistant turn.

    The worker image includes Tesseract and the API uses the same language
    defaults. If OCR is not installed in a development environment, the
    attachment remains usable as a visual preview and the response explains
    the missing capability instead of pretending the image was read.
    """
    if os.getenv("DOCPILOT_ASSISTANT_OCR_ENABLED", "true").lower() in {"0", "false", "off", "no"}:
        return _response(
            name=name,
            kind="image",
            mime_type=mime_type,
            size=len(data),
            status="unsupported",
            text="",
            error="图片 OCR 已被配置关闭。",
        )

    binary = shutil.which("tesseract")
    if not binary:
        return _response(
            name=name,
            kind="image",
            mime_type=mime_type,
            size=len(data),
            status="unsupported",
            text="",
            error="当前运行环境未安装 OCR 引擎；图片仍可预览，但暂不能提取文字。",
        )

    languages = os.getenv("DOCPILOT_ASSISTANT_OCR_LANG", "chi_sim+eng")
    try:
        result = subprocess.run(
            [binary, "stdin", "stdout", "-l", languages, "--psm", "6"],
            input=data,
            capture_output=True,
            timeout=20,
            check=False,
        )
        if result.returncode != 0:
            raise RuntimeError("tesseract exited with a non-zero status")
        text = _normalize_text(result.stdout.decode("utf-8", errors="replace"))
    except Exception as exc:  # Defensive: a bad image must not break chat.
        logger.warning("Assistant image OCR failed: filename=%s error_type=%s", name, type(exc).__name__)
        return _response(
            name=name,
            kind="image",
            mime_type=mime_type,
            size=len(data),
            status="failed",
            text="",
            error="图片 OCR 失败，请确认图片清晰后重试。",
        )

    if not text:
        return _response(
            name=name,
            kind="image",
            mime_type=mime_type,
            size=len(data),
            status="empty",
            text="",
            error="图片中没有识别到可读文字。",
        )
    return _response(
        name=name,
        kind="image",
        mime_type=mime_type,
        size=len(data),
        status="extracted",
        text=_truncate(text, MAX_ATTACHMENT_TEXT_CHARS),
    )


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _normalize_text(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    return "\n".join(lines).strip()


def _truncate(text: str, max_chars: int) -> str:
    if len(text) <= max_chars:
        return text
    return f"{text[:max_chars].rstrip()}\n\n[附件正文已截断]"


def _status_reason(status: str | None) -> str:
    if status == "empty":
        return "文件没有提取到文本。"
    if status == "failed":
        return "解析失败。"
    if status == "unsupported":
        return "格式暂不支持。"
    return "附件没有提供提取文本。"


def _upload_response_from_record(attachment: AssistantAttachment) -> AssistantAttachmentUploadResponse:
    return AssistantAttachmentUploadResponse(
        id=attachment.id,
        name=attachment.original_filename,
        kind=attachment.kind,  # type: ignore[arg-type]
        mime_type=attachment.mime_type,
        size=attachment.size,
        extraction_status=attachment.extraction_status,  # type: ignore[arg-type]
        extracted_text=attachment.extracted_text,
        error=attachment.extraction_error,
    )


def _payload_from_record(attachment: AssistantAttachment) -> AssistantAttachmentPayload:
    return AssistantAttachmentPayload(
        id=attachment.id,
        name=attachment.original_filename,
        kind=attachment.kind,  # type: ignore[arg-type]
        mime_type=attachment.mime_type,
        size=attachment.size,
        extraction_status=attachment.extraction_status,  # type: ignore[arg-type]
        extracted_text=attachment.extracted_text,
        document_id=attachment.document_id,
        error=attachment.extraction_error,
    )


def _load_staged_attachments_for_ingestion(
    db: Session,
    *,
    current_user: CurrentUser,
    attachment_ids: list[str],
) -> list[AssistantAttachment]:
    records = {
        record.id: record
        for record in db.scalars(
            select(AssistantAttachment)
            .where(
                AssistantAttachment.id.in_(attachment_ids),
                AssistantAttachment.user_id == current_user.id,
                AssistantAttachment.org_id == current_user.org_id,
            )
            .with_for_update()
        )
    }
    if len(records) != len(attachment_ids):
        raise HTTPException(status_code=404, detail="Assistant attachment not found")
    ordered = [records[attachment_id] for attachment_id in attachment_ids]
    for attachment in ordered:
        _require_staged_attachment(attachment)
    return ordered


def _require_staged_attachment(attachment: AssistantAttachment) -> None:
    if attachment.expires_at <= _utcnow():
        attachment.status = "expired"
        raise HTTPException(status_code=410, detail="Assistant attachment expired; upload it again")
    if attachment.status != "staged":
        raise HTTPException(status_code=409, detail="Assistant attachment is no longer available for ingestion")


def _find_reusable_agent_bundle(db: Session, project_id: str, label: str) -> Bundle | None:
    return db.scalar(
        select(Bundle)
        .where(
            Bundle.project_id == project_id,
            Bundle.label == label,
            Bundle.source_type == "assistant_upload",
            Bundle.ingest_status.not_in(("queued", "running", "indexing")),
        )
        .order_by(Bundle.created_at.desc())
    )


def _queue_bundle_ingestion(db: Session, bundle: Bundle, current_user: CurrentUser) -> bool:
    try:
        celery.send_task("worker.ingest_bundle", args=[bundle.id])
        return True
    except Exception as exc:  # Keep durable documents available for manual retry.
        logger.warning("Assistant attachment ingestion dispatch failed: %s", type(exc).__name__)
        bundle.ingest_status = "ready_to_ingest"
        record_audit_event(
            db,
            project_id=bundle.project_id,
            event_type="assistant.attachment_ingest_dispatch_failed",
            actor_type="assistant",
            actor_id=current_user.id,
            payload={"bundle_id": bundle.id},
        )
        db.commit()
        return False


def _delete_storage_keys_quietly(storage_keys: list[str]) -> None:
    if not storage_keys:
        return
    try:
        from app.adapters.storage import delete_storage_key

        for storage_key in storage_keys:
            delete_storage_key(storage_key)
    except Exception as exc:
        logger.warning("Assistant attachment staged-object cleanup failed: %s", type(exc).__name__)


def _delete_attached_staging_objects(db: Session, attachments: list[AssistantAttachment]) -> None:
    deleted = False
    for attachment in attachments:
        if not attachment.storage_key:
            continue
        try:
            from app.adapters.storage import delete_storage_key

            delete_storage_key(attachment.storage_key)
        except Exception as exc:
            logger.warning("Assistant attachment staged-object cleanup failed: %s", type(exc).__name__)
            continue
        attachment.storage_key = ""
        deleted = True
    if not deleted:
        return
    try:
        db.commit()
    except Exception as exc:
        db.rollback()
        logger.warning("Assistant attachment cleanup state persistence failed: %s", type(exc).__name__)


def _safe_filename(filename: str) -> str:
    value = filename.replace("\\", "/").rsplit("/", 1)[-1].strip()
    return value[:240] or "untitled"


def _utcnow() -> datetime:
    return datetime.now(UTC).replace(tzinfo=None)
