from io import BytesIO
import json

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .schemas import BidReadinessSummary


def render_readiness_xlsx(summary: BidReadinessSummary) -> bytes:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet["A1"] = "BidPilot Bid Readiness Pack"
    summary_sheet["A1"].font = Font(size=18, bold=True, color="17324D")
    summary_sheet["A3"] = "Project"
    summary_sheet["B3"] = summary.project_name
    summary_sheet["A4"] = "Response readiness"
    summary_sheet["B4"] = summary.readiness_score
    summary_sheet["A5"] = "Formula version"
    summary_sheet["B5"] = summary.formula_version
    summary_sheet["A6"] = "Source fingerprint"
    summary_sheet["B6"] = summary.source_fingerprint

    score_rows = [
        ("Mandatory closure", summary.scores.mandatory_closure),
        ("Scored coverage", summary.scores.scored_coverage),
        ("Verification", summary.scores.verification),
        ("Assignment", summary.scores.assignment),
    ]
    for row_index, (label, value) in enumerate(score_rows, start=8):
        summary_sheet.cell(row=row_index, column=1, value=label)
        summary_sheet.cell(row=row_index, column=2, value=value)
        summary_sheet.cell(row=row_index, column=2).number_format = "0.0%"
    summary_sheet.column_dimensions["A"].width = 24
    summary_sheet.column_dimensions["B"].width = 72

    matrix = workbook.create_sheet("Requirements")
    headers = [
        "ID",
        "Section",
        "Requirement",
        "Category",
        "Mandatory",
        "Score weight",
        "Risk",
        "Coverage",
        "Evidence",
        "Verification",
        "Owner",
        "Reviewer",
        "Due",
        "Source locator",
    ]
    matrix.append(headers)
    header_fill = PatternFill("solid", fgColor="17324D")
    for cell in matrix[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for requirement in summary.requirements:
        matrix.append(
            [
                requirement.id,
                requirement.section_key,
                requirement.requirement_text,
                requirement.bid_category,
                "Yes" if requirement.is_mandatory else "No",
                requirement.score_weight,
                requirement.risk_level,
                requirement.coverage_status,
                requirement.evidence_status,
                requirement.verification_status,
                requirement.owner_user_id,
                requirement.reviewer_user_id,
                requirement.due_at.isoformat() if requirement.due_at else None,
                json.dumps(requirement.source_locator_json, ensure_ascii=False)
                if requirement.source_locator_json
                else None,
            ]
        )
    widths = [38, 20, 64, 18, 12, 14, 12, 16, 16, 16, 38, 38, 22, 48]
    for index, width in enumerate(widths, start=1):
        matrix.column_dimensions[matrix.cell(row=1, column=index).column_letter].width = width
    matrix.freeze_panes = "A2"
    matrix.auto_filter.ref = matrix.dimensions

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_readiness_docx(summary: BidReadinessSummary) -> bytes:
    document = Document()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("BidPilot Bid Readiness Pack")
    run.bold = True
    run.font.size = Pt(20)

    document.add_heading(summary.project_name, level=1)
    document.add_paragraph(
        f"Response readiness: {summary.readiness_score:.1f}/100 "
        f"(formula {summary.formula_version}; this is not a win probability)."
    )
    document.add_paragraph(f"Source fingerprint: {summary.source_fingerprint}")

    document.add_heading("Executive summary", level=2)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Count / score"
    table.rows[0].cells[2].text = "Meaning"
    summary_rows = [
        ("Mandatory closure", f"{summary.scores.mandatory_closure:.1%}", "Mandatory items closed"),
        ("Scored coverage", f"{summary.scores.scored_coverage:.1%}", "Weighted scoring opportunity covered"),
        ("Verification", f"{summary.scores.verification:.1%}", "Requirements human-verified"),
        ("Assignment", f"{summary.scores.assignment:.1%}", "Requirements with an owner"),
        ("Mandatory gaps", str(len(summary.mandatory_gaps)), "Blocking or incomplete mandatory items"),
        ("Evidence gaps", str(len(summary.evidence_gaps)), "Missing, weak, or conflicting evidence"),
    ]
    for label, value, meaning in summary_rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        cells[2].text = meaning

    _add_requirement_section(document, "Mandatory gaps", summary.mandatory_gaps)
    _add_requirement_section(document, "Evidence gaps", summary.evidence_gaps)
    _add_requirement_section(document, "Contradictions", summary.contradictions)
    _add_requirement_section(document, "Overdue ownership", summary.overdue)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_requirement_section(document: Document, title: str, requirements: list) -> None:
    document.add_heading(title, level=2)
    if not requirements:
        document.add_paragraph("None")
        return
    for requirement in requirements:
        document.add_paragraph(
            f"[{requirement.risk_level}] {requirement.requirement_text} "
            f"({requirement.coverage_status}/{requirement.evidence_status})",
            style="List Bullet",
        )


__all__ = ["render_readiness_docx", "render_readiness_xlsx"]
