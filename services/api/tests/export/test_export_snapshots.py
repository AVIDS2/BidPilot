from io import BytesIO
from uuid import uuid4

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.main import app


def _create_approved_deliverable(client: TestClient, content: str) -> tuple[str, str, str]:
    suffix = uuid4().hex[:10]
    project = client.post(
        "/projects",
        json={"name": f"Export Snapshot {suffix}", "scenario_package": "bidpilot"},
    )
    assert project.status_code == 201, project.text
    deliverable = client.post(
        "/deliverables",
        json={
            "project_id": project.json()["id"],
            "type": "proposal",
            "title": f"Proposal {suffix}",
        },
    )
    assert deliverable.status_code == 201, deliverable.text
    section = client.post(
        "/deliverables/sections",
        json={
            "deliverable_id": deliverable.json()["id"],
            "section_key": "technical-approach",
            "title": "Technical Approach",
        },
    )
    assert section.status_code == 201, section.text

    from app.db import SessionLocal
    from app.models import DeliverableSection, SectionVersion

    db = SessionLocal()
    try:
        persisted_section = db.get(DeliverableSection, section.json()["id"])
        assert persisted_section is not None
        version = SectionVersion(
            deliverable_section_id=persisted_section.id,
            version_number=1,
            content_markdown=content,
            created_by_actor="ai",
        )
        db.add(version)
        db.flush()
        db.commit()
        section_id = persisted_section.id
        version_id = version.id
    finally:
        db.close()

    decision = client.post(
        "/review/decisions",
        json={
            "section_id": section_id,
            "section_version_id": version_id,
            "decision": "approved",
            "comment": "Approved for export snapshot.",
        },
    )
    assert decision.status_code == 201, decision.text
    return deliverable.json()["id"], section_id, version_id


def test_export_snapshot_is_idempotent_and_has_durable_downloads() -> None:
    client = TestClient(app)
    deliverable_id, _section_id, approved_version_id = _create_approved_deliverable(
        client,
        "Approved v1 response body.",
    )
    request_id = f"export-{uuid4()}"

    created = client.post(
        f"/export/deliverables/{deliverable_id}",
        json={"client_request_id": request_id},
    )
    assert created.status_code == 201, created.text
    first = created.json()
    assert first["status"] == "generated"
    assert first["version_number"] == 1
    assert first["docx_sha256"]
    assert first["pdf_sha256"]
    assert first["approved_sections"] == [
        {
            "deliverable_section_id": first["approved_sections"][0]["deliverable_section_id"],
            "section_key": "technical-approach",
            "title": "Technical Approach",
            "section_version_id": approved_version_id,
            "version_number": 1,
            "content_sha256": first["approved_sections"][0]["content_sha256"],
        }
    ]

    replayed = client.post(
        f"/export/deliverables/{deliverable_id}",
        json={"client_request_id": request_id},
    )
    assert replayed.status_code == 201, replayed.text
    assert replayed.json()["id"] == first["id"]

    history = client.get(f"/export/deliverables/{deliverable_id}")
    assert history.status_code == 200, history.text
    assert [record["id"] for record in history.json()] == [first["id"]]

    downloaded = client.get(first["docx_download_path"])
    assert downloaded.status_code == 200, downloaded.text
    text = "\n".join(paragraph.text for paragraph in Document(BytesIO(downloaded.content)).paragraphs)
    assert "Approved v1 response body." in text


def test_export_snapshot_remains_pinned_after_a_new_approved_version() -> None:
    client = TestClient(app)
    deliverable_id, section_id, first_version_id = _create_approved_deliverable(
        client,
        "Approved v1 response body.",
    )
    first = client.post(
        f"/export/deliverables/{deliverable_id}",
        json={"client_request_id": f"first-{uuid4()}"},
    )
    assert first.status_code == 201, first.text
    first_record = first.json()

    from app.db import SessionLocal
    from app.models import DeliverableSection, SectionVersion

    db = SessionLocal()
    try:
        section = db.get(DeliverableSection, section_id)
        assert section is not None
        next_version = SectionVersion(
            deliverable_section_id=section.id,
            version_number=2,
            content_markdown="Approved v2 response body.",
            created_by_actor="ai",
        )
        db.add(next_version)
        db.flush()
        section.status = "approved"
        section.approved_version_id = next_version.id
        db.commit()
    finally:
        db.close()

    second = client.post(
        f"/export/deliverables/{deliverable_id}",
        json={"client_request_id": f"second-{uuid4()}"},
    )
    assert second.status_code == 201, second.text
    second_record = second.json()
    assert second_record["version_number"] == 2
    assert second_record["snapshot_hash"] != first_record["snapshot_hash"]
    assert second_record["approved_sections"][0]["section_version_id"] != first_version_id

    first_detail = client.get(f"/export/records/{first_record['id']}")
    assert first_detail.status_code == 200, first_detail.text
    assert first_detail.json()["approved_sections"][0]["section_version_id"] == first_version_id

    old_docx = client.get(first_record["docx_download_path"])
    assert old_docx.status_code == 200, old_docx.text
    old_text = "\n".join(paragraph.text for paragraph in Document(BytesIO(old_docx.content)).paragraphs)
    assert "Approved v1 response body." in old_text
    assert "Approved v2 response body." not in old_text


def test_export_snapshot_rejects_unapproved_content() -> None:
    client = TestClient(app)
    suffix = uuid4().hex[:10]
    project = client.post(
        "/projects",
        json={"name": f"No Approved Export {suffix}", "scenario_package": "bidpilot"},
    )
    assert project.status_code == 201, project.text
    deliverable = client.post(
        "/deliverables",
        json={
            "project_id": project.json()["id"],
            "type": "proposal",
            "title": f"Proposal {suffix}",
        },
    )
    assert deliverable.status_code == 201, deliverable.text

    response = client.post(
        f"/export/deliverables/{deliverable.json()['id']}",
        json={"client_request_id": f"empty-{uuid4()}"},
    )
    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "approved_content_required"


def test_rejected_new_candidate_preserves_prior_approved_export_snapshot() -> None:
    client = TestClient(app)
    deliverable_id, section_id, approved_version_id = _create_approved_deliverable(
        client,
        "Approved v1 response body.",
    )
    first = client.post(
        f"/export/deliverables/{deliverable_id}",
        json={"client_request_id": f"first-{uuid4()}"},
    )
    assert first.status_code == 201, first.text
    first_record = first.json()

    from app.db import SessionLocal
    from app.models import AuditEvent, Deliverable, DeliverableSection, SectionVersion

    db = SessionLocal()
    try:
        section = db.get(DeliverableSection, section_id)
        assert section is not None
        candidate = SectionVersion(
            deliverable_section_id=section.id,
            version_number=2,
            content_markdown="Rejected v2 response body must never export.",
            created_by_actor="ai",
        )
        db.add(candidate)
        db.commit()
        candidate_id = candidate.id
    finally:
        db.close()

    rejected = client.post(
        "/review/decisions",
        json={
            "section_id": section_id,
            "section_version_id": candidate_id,
            "decision": "rejected",
            "comment": "The new candidate changes the commercial commitment.",
        },
    )
    assert rejected.status_code == 201, rejected.text
    assert rejected.json()["decision"] == "rejected"

    # Rejection of a newer candidate cannot revoke or replace the previously
    # approved snapshot. A retry therefore reuses the original export record.
    replayed = client.post(
        f"/export/deliverables/{deliverable_id}",
        json={"client_request_id": f"after-reject-{uuid4()}"},
    )
    assert replayed.status_code == 201, replayed.text
    assert replayed.json()["id"] == first_record["id"]
    assert replayed.json()["approved_sections"][0]["section_version_id"] == approved_version_id

    docx = client.get(first_record["docx_download_path"])
    assert docx.status_code == 200, docx.text
    exported_text = "\n".join(paragraph.text for paragraph in Document(BytesIO(docx.content)).paragraphs)
    assert "Approved v1 response body." in exported_text
    assert "Rejected v2 response body must never export." not in exported_text

    db = SessionLocal()
    try:
        section = db.get(DeliverableSection, section_id)
        deliverable = db.get(Deliverable, deliverable_id)
        assert section is not None
        assert deliverable is not None
        assert section.approved_version_id == approved_version_id
        assert section.status == "approved"
        assert deliverable.status == "approved"
        event_types = set(
            db.scalars(
                select(AuditEvent.event_type).where(AuditEvent.project_id == deliverable.project_id)
            ).all()
        )
    finally:
        db.close()

    assert {"review.approved", "deliverable.approved", "review.rejected", "deliverable.exported"} <= event_types
